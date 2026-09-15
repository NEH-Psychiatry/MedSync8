"""Tests for scripts/credentialing_alert.py — synthetic data only, no network."""
from __future__ import annotations

import io
import sys
import urllib.error
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))

import credentialing_alert as ca  # noqa: E402

TODAY = date(2026, 6, 1)  # frozen; all fixture dates are relative to this


def _lic(**kw) -> ca.LicenseRow:
    base = dict(license_id=1, employee_id=1, credential_id=1, status_id=1,
                issue_date=None, expiry_date=None, status_name="Active", risk_weight=1)
    base.update(kw)
    lic = ca.LicenseRow(**base)
    if lic.expiry_date is not None:
        lic.days_to_expiry = (lic.expiry_date - TODAY).days
    return lic


# ---------------------------------------------------------------------------
# Date coercion
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value,expected", [
    (None, None), ("", None),
    (datetime(2026, 7, 4, 13, 0), date(2026, 7, 4)),
    (date(2026, 7, 4), date(2026, 7, 4)),
    ("2026-07-04", date(2026, 7, 4)),
    ("2026-07-04T00:00:00", date(2026, 7, 4)),
    ("not a date", None), (12345, None),
])
def test_coerce_date(value, expected):
    assert ca._coerce_date(value) == expected


# ---------------------------------------------------------------------------
# Urgency bands (today is passed in, never read from the clock)
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("days,band", [
    (-1, "LAPSED"), (0, "CRITICAL"), (30, "CRITICAL"), (31, "URGENT"),
    (60, "URGENT"), (61, "WATCH"), (90, "WATCH"), (91, "OK"), (400, "OK"),
])
def test_bands_by_days_to_expiry(days, band):
    lic = _lic(expiry_date=TODAY + timedelta(days=days), status_name="Expiring Soon")
    ca.classify(lic, TODAY)
    assert lic.urgency_band == band
    assert (lic.action_required == "") == (band == "OK")


def test_pending_vs_stalled_uses_passed_today():
    fresh = _lic(issue_date=TODAY - timedelta(days=10))
    old = _lic(issue_date=TODAY - timedelta(days=61))
    ca.classify(fresh, TODAY); ca.classify(old, TODAY)
    assert fresh.urgency_band == "PENDING"
    assert old.urgency_band == "STALLED"
    # A different "today" must change the answer — proves no date.today() inside.
    later = _lic(issue_date=TODAY - timedelta(days=10))
    ca.classify(later, TODAY + timedelta(days=100))
    assert later.urgency_band == "STALLED"


def test_status_mismatch_flag_for_expiring_credential_marked_active():
    lic = _lic(expiry_date=TODAY + timedelta(days=20), status_name="Active")
    ca.classify(lic, TODAY)
    assert "DimStatus shows 'Active'" in lic.action_required
    ok = _lic(expiry_date=TODAY + timedelta(days=20), status_name="Expiring Soon")
    ca.classify(ok, TODAY)
    assert "DimStatus" not in ok.action_required


def test_rank_orders_lapsed_before_critical_then_by_days_then_risk():
    rows = [
        _lic(license_id=1, expiry_date=TODAY + timedelta(days=5), risk_weight=1),
        _lic(license_id=2, expiry_date=TODAY - timedelta(days=3), risk_weight=1),
        _lic(license_id=3, expiry_date=TODAY + timedelta(days=5), risk_weight=5),
    ]
    for r in rows:
        ca.classify(r, TODAY)
    ordered = [r.license_id for r in sorted(rows, key=ca.rank_key)]
    assert ordered == [2, 3, 1]


# ---------------------------------------------------------------------------
# Workbook loading (real openpyxl tables, synthetic values)
# ---------------------------------------------------------------------------

def _table(wb: Workbook, name: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r)
    ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    ws.add_table(Table(displayName=name, ref=ref))


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _table(wb, "FactLicensing",
           ["LicenseID", "EmployeeID", "CredentialID", "StatusID", "IssueDate", "ExpiryDate", "Notes"],
           [[1, 10, 100, 1, datetime(2024, 1, 1), "2026-06-20", None],
            [2, 11, 101, 2, datetime(2026, 3, 1), None, "application open"],
            [3, 99, 100, 1, None, None, None]])  # employee not in dim → fallback name
    _table(wb, "DimEmployee", ["EmployeeID", "FullName", "Role", "SupervisingPhysicianID"],
           [[10, "Test-Alpha, Nurse", "PMHNP", 1], [11, "Test-Beta, PA", "PA", ""],
            [11, "Test-Beta-DUP, PA", "PA", ""]])  # duplicate key → warning, later wins
    _table(wb, "DimCredential", ["CredentialID", "CredentialName", "CredentialType", "Jurisdiction", "IssuingAuthority"],
           [[100, "RN License", "License", "WI", "WI DSPS"], [101, "DEA Registration", "Registration", "US", "DEA"]])
    _table(wb, "DimStatus", ["StatusID", "StatusName", "StatusCategory", "RiskWeight"],
           [[1, "Active", "Current", 3], [2, "Pending", "Application", 1]])
    p = tmp_path / "synthetic-star-schema.xlsx"
    wb.save(p)
    return p


def test_load_workbook_enriches_and_warns_on_duplicate_dim(workbook_path, caplog):
    rows = ca.load_workbook_data(str(workbook_path), TODAY)
    assert [r.license_id for r in rows] == [1, 2, 3]
    r1, r2, r3 = rows
    assert r1.employee_name == "Test-Alpha, Nurse" and r1.jurisdiction == "WI"
    assert r1.days_to_expiry == 19 and r1.risk_weight == 3
    assert r2.expiry_date is None and r2.supervising_id is None  # "" coerced to None
    assert r2.employee_name == "Test-Beta-DUP, PA"  # later duplicate row wins
    assert r3.employee_name == "Employee #99"  # missing dim → fallback
    assert any("Duplicate EmployeeID=11" in m for m in caplog.messages)


def test_missing_table_raises_keyerror(tmp_path):
    wb = Workbook(); wb.save(tmp_path / "empty.xlsx")
    with pytest.raises(KeyError):
        ca.load_workbook_data(str(tmp_path / "empty.xlsx"), TODAY)


# ---------------------------------------------------------------------------
# Teams delivery (urlopen mocked; no network)
# ---------------------------------------------------------------------------

class _Resp:
    def __init__(self, status): self.status = status
    def __enter__(self): return self
    def __exit__(self, *a): return False


def test_card_shape_empty_vs_actionable():
    empty = ca._build_teams_card([], TODAY)
    assert empty["themeColor"] == "00C176" and "sections" not in empty
    lic = _lic(expiry_date=TODAY + timedelta(days=5), status_name="Expiring Soon")
    ca.classify(lic, TODAY)
    card = ca._build_teams_card([lic], TODAY)
    assert card["themeColor"] == "FF4444" and len(card["sections"]) == 1
    assert "Critical" in card["text"]


def test_send_teams_returns_false_when_unconfigured(monkeypatch):
    monkeypatch.setattr(ca, "TEAMS_WEBHOOK_URL", "")
    assert ca.send_teams([], TODAY) is False


def test_send_teams_success_and_failure(monkeypatch):
    monkeypatch.setattr(ca, "TEAMS_WEBHOOK_URL", "https://example.invalid/hook")
    monkeypatch.setattr(ca, "TEAMS_CHANNEL_WEBHOOK", "")
    calls: list[dict] = []

    def fake_urlopen(req, timeout=None):
        calls.append({"url": req.full_url, "timeout": timeout})
        return _Resp(200)

    monkeypatch.setattr(ca.urllib.request, "urlopen", fake_urlopen)
    assert ca.send_teams([], TODAY) is True
    assert calls[0]["timeout"] == 15  # a stalled webhook must never block forever

    def failing_urlopen(req, timeout=None):
        raise urllib.error.URLError("boom")

    monkeypatch.setattr(ca.urllib.request, "urlopen", failing_urlopen)
    assert ca.send_teams([], TODAY) is False


def test_main_exit_code_2_when_workbook_missing(monkeypatch, tmp_path):
    monkeypatch.setattr(ca, "WORKBOOK_PATH", str(tmp_path / "nope.xlsx"))
    assert ca.main() == 2


def test_main_exit_code_1_when_delivery_fails(monkeypatch, workbook_path):
    monkeypatch.setattr(ca, "WORKBOOK_PATH", str(workbook_path))
    monkeypatch.setattr(ca, "TEAMS_WEBHOOK_URL", "")
    assert ca.main() == 1
