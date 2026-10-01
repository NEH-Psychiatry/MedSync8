"""Tests for scripts/rate_workbook_sync.py — the Excel bridge between the practice's
rate-verification workbook and the billing tracker.

The fixture reproduces the real workbook's Rates and Inputs sheet layouts
(title row, header row, label / value / unit / evidence / source / limitation)
with the published CY2026 figures. No PHI: payer fee schedules are public data.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))

import rate_workbook_sync as sync  # noqa: E402
from cocm_time_tracker import PlanningAssumptions, _parse_overrides  # noqa: E402

RATES_ROWS = [
    # label, value, unit, evidence, source, limitation
    ("Medicare WI physician 99214", 129.57, "USD / service", "Recomputed from July 2026 file", "https://www.cms.gov/files/zip/rvu26c-updated-06-30-2026.zip", "Nonfacility, non-QP."),
    ("Medicare WI G2082", 912.05, "USD / 56-mg session", "Recomputed from July 2026 file", "https://www.cms.gov/files/zip/rvu26c-updated-06-30-2026.zip", "Drug-inclusive office bundle."),
    ("Medicare FQHC 99492", "$160.32", "USD / initial month", "CMS 2026 designated rate", "https://www.cms.gov/files/zip/rhc-fqhc-cy-2026-non-facility-payment-rates.zip", "National nonfacility FQHC care coordination."),
    ("Medicare FQHC 99493", 144.96, "USD / subsequent month", "CMS 2026 designated rate", "https://www.cms.gov/files/zip/rhc-fqhc-cy-2026-non-facility-payment-rates.zip", "Separately priced CoCM pathway."),
    ("Medicare FQHC 99494", 61.46, "USD / qualifying unit", "CMS 2026 designated rate", "https://www.cms.gov/files/zip/rhc-fqhc-cy-2026-non-facility-payment-rates.zip", "Qualifying additional work only."),
    ("Medicare WI office 99492", 153.19, "USD / initial month", "Recomputed from July 2026 file", "https://www.cms.gov/files/zip/rvu26c-updated-06-30-2026.zip", "Non-FQHC physician billing practitioner."),
    ("Medicare WI office 99493", 138.71, "USD / subsequent month", "Recomputed from July 2026 file", "https://www.cms.gov/files/zip/rvu26c-updated-06-30-2026.zip", "Separate from consultant reimbursement."),
    ("Medicare WI office 99494", 58.72, "USD / qualifying unit", "Recomputed from July 2026 file", "https://www.cms.gov/files/zip/rvu26c-updated-06-30-2026.zip", "Additional eligible work only."),
    ("Wisconsin 99214 schedule benchmark", 85.58, "USD / service", "Official 2026-09-05 fee snapshot", "https://www.forwardhealth.wi.gov/MaxFee/MEDSV_max_fee_20260905.csv", "Unmodified PT6/physician rows."),
    ("Wisconsin FFS 99492", 146.05, "USD / initial month", "Official 2026-09-05 fee snapshot", "https://www.forwardhealth.wi.gov/MaxFee/MEDSV_max_fee_20260905.csv", "CHCs use the T1015 PPS pathway instead."),
    ("Wisconsin FFS 99493", 141.61, "USD / subsequent month", "Official 2026-09-05 fee snapshot", "https://www.forwardhealth.wi.gov/MaxFee/MEDSV_max_fee_20260905.csv", "Psychiatrist cannot be the Wisconsin CoCM billing practitioner."),
    ("Wisconsin FFS 99494", 60.51, "USD / qualifying unit", "Official 2026-09-05 fee snapshot", "https://www.forwardhealth.wi.gov/MaxFee/MEDSV_max_fee_20260905.csv", "No added PPS encounter assumed for an extra time unit."),
    ("Wisconsin J0013 maximum", 14.90, "USD / mg", "Official 2026-09-05 fee snapshot", "https://www.forwardhealth.wi.gov/MaxFee/MEDSV_max_fee_20260905.csv", "Effective 2026-01-01."),
    ("Wisconsin FQHC GAF", 0.980, "factor", "CMS 2026 GAF file", "https://www.cms.gov/files/zip/2026-fqhc-gafs.zip", "MAC 06302, locality 00."),
]
INPUT_ROWS = [
    ("Collection realization", "94.0%", "fraction of allowance", "Illustrative / editable", "Aggregate payer + patient receipts."),
    ("Physician loaded labor", "$250", "USD / hour", "Illustrative / editable", "Compensation or opportunity cost."),
    ("CoCM initial-month share", 0.15, "fraction of billable months", "Illustrative / editable", "Remaining months use 99493."),
    ("CoCM extra units / billable month", 0.2, "average qualifying units", "Illustrative / editable", "Analyst utilization test."),
    ("CoCM billable conversion", "80.0%", "billable / active patients", "Illustrative / editable", "100 billable months require 125 active patients."),
    ("Active caseload / BHCM FTE", 60, "active patients / FTE", "Illustrative / editable", "AIMS example for complex/FQHC populations."),
    ("BHCM budget increment", 0.1, "FTE", "Illustrative / editable", "Round staffing up to this budget increment."),
]


def build_workbook(path: Path, rates=RATES_ROWS, inputs=INPUT_ROWS) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["NEH psychiatry, Spravato and FQHC economics"])
    inp = wb.create_sheet("Inputs")
    inp.append(["Operating assumptions"])
    inp.append([])
    inp.append(["Assumption", "Value", "Unit", "Evidence status", "Basis / limitation"])
    for r in inputs:
        inp.append(list(r))
    rs = wb.create_sheet("Rates")
    rs.append(["Payment and acquisition evidence"])
    rs.append([])
    rs.append(["Rate / input", "Value", "Unit", "Evidence status", "Source / locator", "Use limitation"])
    for r in rates:
        rs.append(list(r))
    wb.save(path)
    return path


@pytest.fixture
def workbook(tmp_path) -> Path:
    return build_workbook(tmp_path / "NEH-CoCM-Test-Workbook.xlsx")


def test_rates_sheet_parses_only_payer_rate_rows_for_catalogue_codes(workbook):
    rates = sync.load_workbook_rates(workbook)
    got = {(r.payer, r.code): r for r in rates}
    assert set(got) == {(p, c) for p in ("medicare-wi", "medicare-fqhc", "wi-medicaid")
                        for c in ("99492", "99493", "99494")}
    assert got[("medicare-fqhc", "99492")].usd == 160.32          # "$160.32" string cell
    assert got[("wi-medicaid", "99493")].confidence == "verified_primary"
    assert "MEDSV_max_fee_20260905" in got[("wi-medicaid", "99493")].source
    assert "read from NEH-CoCM-Test-Workbook.xlsx" in got[("wi-medicaid", "99493")].source
    assert "billing practitioner" in got[("wi-medicaid", "99493")].note
    # --all keeps non-catalogue payer rows (Spravato bundle, E/M benchmark); rows whose
    # label is not a payer rate ("Wisconsin 99214 schedule benchmark", "Wisconsin
    # J0013 maximum", the GAF factor) are never emitted.
    everything = sync.load_workbook_rates(workbook, include_all=True)
    extra = {(r.payer, r.code) for r in everything} - set(got)
    assert extra == {("medicare-wi", "99214"), ("medicare-wi", "G2082")}
    assert all(not r.in_catalogue for r in everything if (r.payer, r.code) in extra)


def test_reconcile_matches_tracker_catalogue_exactly(workbook):
    rows = sync.reconcile(sync.load_workbook_rates(workbook))
    assert len(rows) == 9 and {r.status for r in rows} == {"match"}
    assert all(r.tracker_confidence == "verified_primary" for r in rows)


def test_reconcile_flags_mismatch_and_unknown_slots(tmp_path):
    rows = list(RATES_ROWS)
    rows[9] = ("Wisconsin FFS 99492", 150.00) + rows[9][2:]                      # changed fee
    rows.append(("Wisconsin FFS 99484", 41.28, "USD / month", "Official 2026-09-05 fee snapshot", "csv", ""))
    wb = build_workbook(tmp_path / "wb.xlsx", rates=rows)
    rec = {(r.payer, r.code): r for r in sync.reconcile(sync.load_workbook_rates(wb, include_all=True))}
    assert rec[("wi-medicaid", "99492")].status == "mismatch" and rec[("wi-medicaid", "99492")].tracker_usd == 146.05
    assert rec[("wi-medicaid", "99484")].status == "not_in_tracker"              # catalogue code, no WI fee yet
    assert rec[("medicare-wi", "G2082")].status == "not_in_catalogue"


def test_evidence_status_mapping_never_upgrades_illustrative_rows():
    assert sync.confidence_for("Official 2026-09-05 fee snapshot") == "verified_primary"
    assert sync.confidence_for("CMS 2026 designated rate") == "verified_primary"
    assert sync.confidence_for("Recomputed from July 2026 file") == "verified_primary"
    assert sync.confidence_for("Memo V1.0 grid") == "verified_secondary"
    assert sync.confidence_for("Illustrative / editable") is None
    assert sync.confidence_for("") is None
    assert sync.classify_label("NADAC 56-mg pack device") is None
    assert sync.classify_label("Medicare WI office 99492") == ("medicare-wi", "99492")
    assert sync.classify_label("Medicare WI G2214") == ("medicare-wi", "G2214")


def test_overrides_file_round_trips_into_the_tracker(workbook, tmp_path):
    out = tmp_path / "rates.json"
    out.write_text(json.dumps(sync.to_overrides(sync.load_workbook_rates(workbook))))
    ov = _parse_overrides({"RATE_OVERRIDES_FILE": str(out)})
    assert ov["wi-medicaid"]["99494"].usd == 60.51
    assert ov["wi-medicaid"]["99494"].confidence == "verified_primary"
    assert "Rates sheet" in ov["wi-medicaid"]["99494"].source
    # Inline JSON still wins over the file; an unreadable file is ignored with a warning.
    both = _parse_overrides({"RATE_OVERRIDES_FILE": str(out),
                             "RATE_OVERRIDES_JSON": '{"wi-medicaid": {"99494": 61.00}}'})
    assert both["wi-medicaid"]["99494"].usd == 61.00
    assert _parse_overrides({"RATE_OVERRIDES_FILE": str(tmp_path / "missing.json")})["wi-medicaid"] == {}


def test_planning_assumptions_read_percentages_and_counts(workbook):
    a = sync.load_planning_assumptions(workbook)
    assert a == PlanningAssumptions(billable_conversion=0.80, initial_month_share=0.15,
                                    extra_units_per_month=0.2, caseload_per_fte=60,
                                    fte_increment=0.1, collection_realization=0.94)


def test_cli_check_and_exports(workbook, tmp_path):
    env = {"PYTHONPATH": str(REPO / "scripts")}
    r = subprocess.run([sys.executable, str(REPO / "scripts" / "rate_workbook_sync.py"),
                        "--workbook", str(workbook), "--check",
                        "--write-overrides", str(tmp_path / "o.json"),
                        "--write-assumptions", str(tmp_path / "a.json")],
                       capture_output=True, text=True, env={**env, "PATH": "/usr/bin:/bin"})
    assert r.returncode == 0, r.stderr
    assert "CHECK OK" in r.stdout and "9 rate rows read" in r.stdout
    assert json.loads((tmp_path / "o.json").read_text())["medicare-fqhc"]["99493"]["usd"] == 144.96
    assert json.loads((tmp_path / "a.json").read_text())["caseload_per_fte"] == 60
    bad = build_workbook(tmp_path / "bad.xlsx", rates=[("Wisconsin FFS 99492", 1.00, "", "Official snapshot", "", "")])
    r = subprocess.run([sys.executable, str(REPO / "scripts" / "rate_workbook_sync.py"),
                        "--workbook", str(bad), "--check"], capture_output=True, text=True,
                       env={**env, "PATH": "/usr/bin:/bin"})
    assert r.returncode == 1 and "CHECK FAILED" in r.stderr


def test_missing_sheet_is_a_clean_error(tmp_path):
    wb = openpyxl.Workbook(); wb.save(tmp_path / "empty.xlsx")
    with pytest.raises(ValueError, match="no 'Rates' sheet"):
        sync.load_workbook_rates(tmp_path / "empty.xlsx")
